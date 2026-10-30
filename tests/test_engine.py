"""Tests fuer die SpeziSmith-Engine mit den echten Vorlagen."""

import os
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from spezismith.engine import (
    Rule,
    analyze_template,
    apply_to_file,
    detect_language,
    process_batch,
    unique_target_path,
)

TPL_DE = os.path.join(os.path.dirname(__file__), "..", "templates", "vorlage_de.xlsx")
TPL_EN = os.path.join(os.path.dirname(__file__), "..", "templates", "vorlage_en.xlsx")


# ---------------------------------------------------------------------------
# Analyse & Spracherkennung
# ---------------------------------------------------------------------------

def test_detect_language():
    assert detect_language(TPL_DE) == "de"
    assert detect_language(TPL_EN) == "en"


def test_analyze_de_fields():
    fields = analyze_template(TPL_DE)
    by_coord = {f.coord: f for f in fields}

    # Kernfelder vorhanden
    assert "A1" in by_coord          # Firmenkopf
    assert "A5" in by_coord          # Artikelnummer (Beschriftung)
    assert "J1" in by_coord          # "Spezifikation gültig ab ..."
    assert "A44" in by_coord         # Allergene
    assert "A86" in by_coord         # maschinell erstellt

    assert "AULBACH GmbH" in by_coord["A1"].value
    assert by_coord["A5"].value == "Artikelnummer:"

    # leere Eingabefelder neben Beschriftungen wurden erkannt
    inputs_row5 = [f for f in fields if f.is_input and f.row == 5]
    assert inputs_row5, "Eingabefeld neben 'Artikelnummer:' fehlt"

    # keine doppelten Koordinaten
    assert len(by_coord) == len(fields)


def test_analyze_en_fields():
    fields = analyze_template(TPL_EN)
    by_coord = {f.coord: f for f in fields}
    assert by_coord["A3"].value == "Product-Specification"
    assert "B71" in by_coord                      # tare weight Wert 1.15
    assert by_coord["B71"].value.startswith("1.15")


# ---------------------------------------------------------------------------
# Ersetzen (ZIP-Surgery)
# ---------------------------------------------------------------------------

def _sheet_xml(path):
    import re as _re
    import xml.dom.minidom as _md
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    # strikte Wohlgeformtheit (expat bricht bei doppelten Attributen ab)
    _md.parseString(xml.encode("utf-8"))
    body = xml.split("?>", 1)[1]
    root_tag = _re.search(r"<((?:\w+:)?\w+)((?:\s+[^<>]*)?)>", body).group(0)
    # mc:Ignorable-Praefixe muessen deklariert sein (sonst Excel-Reparatur!)
    ign = _re.search(r'(?:\w+:)?Ignorable="([^"]+)"', root_tag)
    if ign:
        for prefix in ign.group(1).split():
            assert f'xmlns:{prefix}=' in root_tag, \
                f"Praefix '{prefix}' in Ignorable nicht deklariert"
    return xml, root_tag


def test_output_is_excel_valid(tmp_path):
    """Beide Vorlagen muessen nach der Bearbeitung Excel-konformes XML haben."""
    for tpl in (TPL_DE, TPL_EN):
        out = tmp_path / (os.path.basename(tpl))
        res = apply_to_file(tpl, str(out), [
            Rule(coord="A1", new="Neue Firma GmbH"),
            Rule(coord="J1", new="01.01.2027", mode="replace", find="xx.xx.xxxx"),
        ])
        assert res.ok, res.error
        xml, root_tag = _sheet_xml(str(out))
        # urspruengliche Namespace-Praefixe bleiben erhalten (kein ns0/ns1...)
        assert "ns0" not in root_tag and "ns1:" not in root_tag


def test_merged_target_written_to_anchor(tmp_path):
    """A1 ist Anker eines Verbundes - Schreiben darf den Verbund nicht zerstoeren."""
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [Rule(coord="B1", new="X")])
    assert res.ok, res.error
    assert res.cells[0].coord == "A1"  # auf Anker gemappt
    wb = openpyxl.load_workbook(out)
    try:
        assert "X" in str(wb.worksheets[0]["A1"].value)
    finally:
        wb.close()


def _read_cell(path, coord):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return wb.worksheets[0][coord].value
    finally:
        wb.close()


def test_apply_full_cell_replacement(tmp_path):
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [
        Rule(coord="A1", new="Test1"),
        Rule(coord="A5", new="Artikelnummer (NEU):"),
    ])
    assert res.ok, res.error
    assert res.language == "de"
    assert res.changed_count == 2
    assert "Test1" in _read_cell(str(out), "A1")
    assert _read_cell(str(out), "A5") == "Artikelnummer (NEU):"


def test_logo_and_parts_preserved(tmp_path):
    """Logo, VML, ActiveX etc. bleiben byte-identisch erhalten."""
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [Rule(coord="A1", new="Test1")])
    assert res.ok, res.error

    with zipfile.ZipFile(TPL_DE) as a, zipfile.ZipFile(out) as b:
        names_a, names_b = set(a.namelist()), set(b.namelist())
        assert names_a == names_b, "ZIP-Struktur veraendert!"
        for name in names_a:
            if name == "xl/worksheets/sheet1.xml":
                continue
            assert a.read(name) == b.read(name), f"Part geaendert: {name}"
        # Logo explizit pruefen
        assert a.read("xl/media/image1.jpeg") == b.read("xl/media/image1.jpeg")


def test_merges_and_styles_survive(tmp_path):
    out = tmp_path / "out.xlsx"
    apply_to_file(TPL_DE, str(out), [Rule(coord="A1", new="Test1")])
    wb = openpyxl.load_workbook(out)
    try:
        ws = wb.worksheets[0]
        assert len(ws.merged_cells.ranges) >= 140
        assert ws["A1"].font is not None  # Stilinformation vorhanden
    finally:
        wb.close()


def test_replace_substring_mode(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="J1", new="01.01.2027", mode="replace", find="xx.xx.xxxx")
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok, res.error
    assert res.changed_count == 1
    assert "01.01.2027" in _read_cell(str(out), "J1")
    assert "xx.xx.xxxx" not in str(_read_cell(str(out), "J1"))


def test_replace_substring_not_found(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="J1", new="X", mode="replace", find="GIBTS-NICHT")
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok
    assert res.cells[0].status == "uebersprungen"


def test_verify_mismatch_skips(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="A1", new="Test1", expected="völlig anderer Wert")
    res = apply_to_file(TPL_DE, str(out), [rule], verify=True)
    assert res.ok
    assert res.cells[0].status == "uebersprungen"
    assert "AULBACH" in _read_cell(str(out), "A1")


def test_numeric_value_stays_numeric(tmp_path):
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_EN, str(out), [Rule(coord="B71", new="2,30")])
    assert res.ok, res.error
    value = _read_cell(str(out), "B71")
    assert value == 2.3  # als Zahl gespeichert, Komma wurde konvertiert


def test_write_into_empty_cell(tmp_path):
    """Leere Eingabefelder (z. B. Wert neben 'Artikelnummer:') fuellen."""
    fields = analyze_template(TPL_DE)
    target_field = next(f for f in fields if f.is_input and f.row == 5)
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [
        Rule(coord=target_field.coord, new="ART-12345"),
    ])
    assert res.ok, res.error
    assert _read_cell(str(out), target_field.coord) == "ART-12345"


def test_clear_cell(tmp_path):
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_EN, str(out), [Rule(coord="B71", new="")])
    assert res.ok, res.error
    assert _read_cell(str(out), "B71") in (None, "")


# ---------------------------------------------------------------------------
# Batch & Zielpfade
# ---------------------------------------------------------------------------

def test_unique_target_path(tmp_path):
    first = unique_target_path(str(tmp_path), "Spezi.xlsx", "_neu")
    assert first.endswith("Spezi_neu.xlsx")
    open(first, "wb").close()
    second = unique_target_path(str(tmp_path), "Spezi.xlsx", "_neu")
    assert "(1)" in second


def test_process_batch_five_files(tmp_path):
    """Fuenf Spezifikationen gleichzeitig verarbeiten."""
    sources = []
    for i in range(5):
        src = tmp_path / f"Spezi_{i}.xlsx"
        with open(TPL_DE, "rb") as fsrc, open(src, "wb") as fdst:
            fdst.write(fsrc.read())
        sources.append(str(src))

    out_dir = tmp_path / "out"
    events = []
    results = process_batch(
        sources,
        [Rule(coord="A1", new="Sammel-Test")],
        str(out_dir),
        suffix="_neu",
        progress=lambda i, n, r: events.append((i, n)),
    )
    assert len(results) == 5
    assert all(r.ok for r in results)
    assert events == [(1, 5), (2, 5), (3, 5), (4, 5), (5, 5)]
    for r in results:
        assert "Sammel-Test" in _read_cell(r.target, "A1")


def test_invalid_file_reports_error(tmp_path):
    bad = tmp_path / "kaputt.xlsx"
    bad.write_text("kein Excel")
    res = apply_to_file(str(bad), str(tmp_path / "o.xlsx"), [Rule(coord="A1", new="X")])
    assert not res.ok
    assert res.error


# ---------------------------------------------------------------------------
# Verschiebungstoleranz, Kopf-/Fußzeile, Layout-Prüfung
# ---------------------------------------------------------------------------

# Simuliert eine Datei, bei der durch 2 Extra-Zutatenzeilen alles ab Zeile 14
# um 2 Zeilen nach unten gerutscht ist ('Vakuum' C78 -> C80 usw.)
SHIFTED_TEXTS = {
    "A3": "Produktspezifizierung",
    "A21": "Sensorik:",            # in der Vorlage: A19
    "C78": "Anzahl Beutel je Karton",  # in der Vorlage: C76
    "C80": "Vakuum",               # in der Vorlage: C78
    "E80": "Atmosphäre",
    "A1": "AULBACH GmbH & Co. KG, Industriestraße 7, 63768 Hösbach",
}


def test_resolve_target_anchor_shifted():
    """'Vakuum'-Fall: Zelle hinter 'Vakuum' (D78) muss im verschobenen Layout
    auf D80 zeigen - nicht auf 'Anzahl Beutel je Karton'."""
    from spezismith.engine import _resolve_target
    rule = Rule(coord="D78", new="JA", expected="",
                anchor_text="Vakuum", anchor_coord="C78")
    coord, via = _resolve_target(rule, SHIFTED_TEXTS, {})
    assert coord == "D80"
    assert via == "anker"


def test_resolve_target_content_shifted():
    """Textzellen werden über ihren Inhalt gefunden, egal wohin sie rutschten."""
    from spezismith.engine import _resolve_target
    rule = Rule(coord="A19", new="Neue Sensorik", expected="Sensorik:")
    coord, via = _resolve_target(rule, SHIFTED_TEXTS, {})
    assert coord == "A21"
    assert via == "inhalt"


def test_resolve_target_fallback_coordinate():
    from spezismith.engine import _resolve_target
    rule = Rule(coord="D78", new="JA")  # ohne expected/anchor
    coord, via = _resolve_target(rule, SHIFTED_TEXTS, {})
    assert coord == "D78"
    assert via == "koordinate"


def test_anchor_rule_on_real_template(tmp_path):
    """Integration auf der echten Vorlage: D78 hinter 'Vakuum' beschreiben,
    'Anzahl Beutel je Karton' (C76) bleibt unberuehrt."""
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="D78", new="JA", expected="",
                anchor_text="Vakuum", anchor_coord="C78")
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok, res.error
    assert res.cells[0].coord == "D78"
    assert _read_cell(str(out), "D78") == "JA"
    assert _read_cell(str(out), "C76") == "Anzahl Beutel je Karton"


def test_layout_similarity():
    from spezismith.engine import layout_fingerprints, layout_similarity
    ref = layout_fingerprints(TPL_DE)
    assert layout_similarity(ref, layout_fingerprints(TPL_DE)) == 1.0
    en = layout_fingerprints(TPL_EN)
    assert layout_similarity(ref, en) < 0.6  # andere Sprache = anderes Layout


def test_verify_empty_expected_requires_empty_target(tmp_path):
    """Verify + leerer Vorlagenwert = nur leere Zielzellen befuellen."""
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="C76", new="X", expected="")  # C76 ist belegt!
    res = apply_to_file(TPL_DE, str(out), [rule], verify=True)
    assert res.ok
    assert res.cells[0].status == "uebersprungen"
    assert "nicht leer" in res.cells[0].message
    assert _read_cell(str(out), "C76") == "Anzahl Beutel je Karton"


def test_header_footer_read():
    from spezismith.engine import get_header_footer
    hf = get_header_footer(TPL_DE)
    assert "HF:oddHeader.L" in hf
    assert "FB_QM_005" in hf["HF:oddHeader.L"]
    assert "&P" in hf["HF:oddHeader.L"]          # Feldcode bleibt sichtbar
    assert "Datum 10.06.2026" in hf["HF:oddFooter.R"]


def test_header_footer_replace_preserves_codes(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="HF:oddFooter.R",
                new="Datum 23.10.2026\n\n\nErsetzt 10.06.2026")
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok, res.error
    assert res.cells[0].status == "geaendert"

    import openpyxl as _opx
    wb = _opx.load_workbook(out)
    try:
        text = wb.worksheets[0].oddFooter.right.text
        assert "23.10.2026" in text
        assert "10.06.2026" in text
        # Seitenzahlen im Kopf unangetastet
        assert "&P" in (wb.worksheets[0].oddHeader.left.text or "")
    finally:
        wb.close()
    # XML-Rohstring: Formatcodes der Fußzeile weiterhin vorhanden
    xml = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "Arial Narrow" in xml
    assert "23.10.2026" in xml


def test_header_footer_missing_part(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="HF:evenFooter.R", new="X")
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok
    assert res.cells[0].status == "uebersprungen"


# ---------------------------------------------------------------------------
# Spezifikation gegen hochgeladene Vorlage angleichen (fehlende Felder)
# ---------------------------------------------------------------------------

def _mini_spec(path, rows, merge=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, (a, b) in enumerate(rows, start=1):
        ws.cell(r, 1, a)
        if b is not None:
            ws.cell(r, 2, b)
    for ref in merge or []:
        ws.merge_cells(ref)
    wb.save(path)


def test_missing_field_is_inserted_without_cell_jumps(tmp_path):
    spec = tmp_path / "spec.xlsx"
    _mini_spec(spec, [("Titel", None), ("Artikel:", "4711"), ("Ende", None)],
               merge=["B3:C3"])
    out = tmp_path / "out.xlsx"
    rules = [
        Rule(coord="B2", new="003", expected="003", label="Revision",
             anchor_text="Revision:", anchor_coord="A2",
             context_above=["Titel"], allow_insert=True),
        Rule(coord="B3", new="Neu", expected="4711", label="Artikel",
             anchor_text="Artikel:", anchor_coord="A3",
             context_above=["Revision:", "Titel"], allow_insert=True),
    ]
    res = apply_to_file(str(spec), str(out), rules)
    assert res.ok, res.error
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(r, 1).value for r in range(1, 5)] == [
        "Titel", "Revision:", "Artikel:", "Ende"]
    assert ws["B2"].value == "003"
    assert ws["B3"].value == "Neu"          # nicht in die falsche Zeile gesprungen
    assert [r.coord for r in ws.merged_cells.ranges] == ["B4:C4"]
    assert any("eingefügt" in c.message for c in res.cells)


def test_missing_field_without_position_is_skipped(tmp_path):
    spec = tmp_path / "spec.xlsx"
    _mini_spec(spec, [("Titel", None), ("Artikel:", "4711")])
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="B2", new="003", expected="003",
                anchor_text="Revision:", anchor_coord="A2",
                context_above=["Gibt es nicht"], allow_insert=True)
    res = apply_to_file(str(spec), str(out), [rule])
    assert res.ok
    assert res.cells[0].status == "uebersprungen"
    ws = openpyxl.load_workbook(out).active
    assert ws["B2"].value == "4711"


def test_multiline_cell_with_other_date_and_initials(tmp_path):
    spec = tmp_path / "spec.xlsx"
    _mini_spec(spec, [("Titel", None),
                      ("Spezifikation gültig ab 01.03.2024\nAB\n", None),
                      ("Ende", None)])
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="A2", new="26.10.2026\nXY", mode="replace",
                find="Spezifikation gültig ab xx.xx.xxxx\nKürzel",
                expected="Spezifikation gültig ab xx.xx.xxxx\nKürzel\n",
                line_start=0, line_count=2,
                context_above=["Titel"], allow_insert=True)
    res = apply_to_file(str(spec), str(out), [rule])
    assert res.ok and res.cells[0].status == "geaendert", res.cells
    ws = openpyxl.load_workbook(out).active
    assert ws["A2"].value == "26.10.2026\nXY\n"
    assert ws["A3"].value == "Ende"


# ---------------------------------------------------------------------------
# Zeilenstruktur der Vorlage uebernehmen
# ---------------------------------------------------------------------------

def _modified_template(tmp_path):
    """Vorlage wie vom Anwender: 'Bestrahlung' -> 'Test 1', plus Zeile 'Test'."""
    wb = openpyxl.load_workbook(TPL_DE)
    ws = wb.active
    ws["A43"] = "Test 1"
    ws["D43"] = "Wert Test 1"
    ws.insert_rows(44)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 44:
            ws.merged_cells.remove(rng)
            ws.merge_cells(start_row=rng.min_row + 1, end_row=rng.max_row + 1,
                           start_column=rng.min_col, end_column=rng.max_col)
    ws["A44"] = "Test"
    ws["D44"] = "Wert Test"
    ws.merge_cells("A44:C44")
    ws.merge_cells("D44:K44")
    path = tmp_path / "tpl.xlsx"
    wb.save(path)
    return str(path)


def _checkbox_rows(path):
    """Zeilenanker (0-basiert) aller Kontrollkaestchen aus den drei Quellen."""
    import re as _re
    with zipfile.ZipFile(path) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
        drawing = z.read("xl/drawings/drawing1.xml").decode("utf-8")
        vml = z.read("xl/drawings/vmlDrawing1.vml").decode("utf-8")
    pair = r"<xdr:col>(\d+)</xdr:col><xdr:colOff>\d+</xdr:colOff><xdr:row>(\d+)</xdr:row>"
    vml_anchors = [a.split(",") for a in
                   _re.findall(r"<x:Anchor>\s*([^<]*?)\s*</x:Anchor>", vml)]
    return {
        "sheet": [(int(c), int(r)) for c, r in _re.findall(pair, sheet)],
        "drawing": [(int(c), int(r)) for c, r in _re.findall(pair, drawing)],
        "vml": [(int(p[0]), int(p[2])) for p in vml_anchors],
    }


def _assert_checkboxes_follow_rows(source, result):
    """Jedes Kaestchen muss in der Zeile mit derselben Beschriftung bleiben."""
    def labels(path):
        ws = openpyxl.load_workbook(path).active
        return {r: str(ws.cell(r, 1).value or "").strip()[:25]
                for r in range(1, ws.max_row + 2)}

    src_labels, res_labels = labels(source), labels(result)
    before, after = _checkbox_rows(source), _checkbox_rows(result)
    for kind, anchors in before.items():
        assert len(anchors) > 0 and len(anchors) == len(after[kind]), kind
        for (col, row), (new_col, new_row) in zip(anchors, after[kind]):
            assert new_col == col
            assert src_labels[row + 1] == res_labels[new_row + 1], (
                f"{kind}: Kaestchen von '{src_labels[row + 1]}' (Zeile {row + 1}) "
                f"steht jetzt bei '{res_labels[new_row + 1]}' (Zeile {new_row + 1})")


def test_structure_sync_replaces_and_inserts_rows(tmp_path):
    tpl = _modified_template(tmp_path)
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [], template=tpl)
    assert res.ok, res.error
    assert res.structure_notes
    ws = openpyxl.load_workbook(out).active
    assert ws["A42"].value == "GVO-Status:"
    assert ws["A43"].value == "Test 1" and ws["D43"].value == "Wert Test 1"
    assert ws["A44"].value == "Test" and ws["D44"].value == "Wert Test"
    assert ws["A45"].value == "Allergene:"
    assert ws["A87"].value.startswith("Dieses Dokument")
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A44:C44", "D44:K44", "A45:A46", "B45:K46"} <= merged
    assert "A43:C43" in merged
    _assert_checkboxes_follow_rows(TPL_DE, out)


def test_structure_sync_keeps_product_values(tmp_path):
    tpl = _modified_template(tmp_path)
    wb = openpyxl.load_workbook(TPL_DE)
    wb.active["C5"] = "Artikel 4711"
    spec = tmp_path / "spec.xlsx"
    wb.save(spec)
    out = tmp_path / "out.xlsx"
    res = apply_to_file(str(spec), str(out), [], template=tpl)
    assert res.ok, res.error
    assert openpyxl.load_workbook(out).active["C5"].value == "Artikel 4711"


def test_structure_sync_identical_template_changes_nothing(tmp_path):
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [], template=TPL_DE)
    assert res.ok and not res.structure_notes


def test_rules_below_inserted_row_keep_their_cells(tmp_path):
    spec = tmp_path / "spec.xlsx"
    _mini_spec(spec, [("Titel", None), ("Ende:", None), ("Letzte", None)])
    out = tmp_path / "out.xlsx"
    # Vorlage hat zusaetzlich Zeile 2 ("Neu:"); Regeln stehen in Vorlagen-Koordinaten
    rules = [
        Rule(coord="B4", new="unten", expected="", label="Letzte"),
        Rule(coord="B2", new="oben", expected="", label="Neu",
             anchor_text="Neu:", anchor_coord="A2",
             context_above=["Titel"], allow_insert=True),
    ]
    res = apply_to_file(str(spec), str(out), rules)
    assert res.ok, res.error
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(r, 1).value for r in range(1, 5)] == [
        "Titel", "Neu:", "Ende:", "Letzte"]
    assert ws["B2"].value == "oben"
    assert ws["B3"].value is None
    assert ws["B4"].value == "unten"


def _insert_template_row(ws, at, label):
    ws.insert_rows(at)
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= at:
            ws.merged_cells.remove(rng)
            ws.merge_cells(start_row=rng.min_row + 1, end_row=rng.max_row + 1,
                           start_column=rng.min_col, end_column=rng.max_col)
    ws.cell(at, 1, label)
    ws.cell(at, 4, "x")
    ws.merge_cells(start_row=at, end_row=at, start_column=1, end_column=3)
    ws.merge_cells(start_row=at, end_row=at, start_column=4, end_column=11)


@pytest.mark.parametrize("insert_at", [[44], [9], [70], [44, 20]])
def test_checkboxes_follow_inserted_rows(tmp_path, insert_at):
    wb = openpyxl.load_workbook(TPL_DE)
    for n, at in enumerate(insert_at):
        _insert_template_row(wb.active, at, f"Neue Zeile {n}")
    tpl = tmp_path / "tpl.xlsx"
    wb.save(tpl)
    out = tmp_path / "out.xlsx"
    res = apply_to_file(TPL_DE, str(out), [], template=str(tpl))
    assert res.ok, res.error
    assert len(res.structure_notes) == len(insert_at)
    _assert_checkboxes_follow_rows(TPL_DE, out)


def test_checkboxes_follow_rule_inserted_row(tmp_path):
    out = tmp_path / "out.xlsx"
    rule = Rule(coord="D44", new="NEU", expected="x", label="Test",
                anchor_text="Test", anchor_coord="A44",
                context_above=["GVO-Status:"], allow_insert=True)
    res = apply_to_file(TPL_DE, str(out), [rule])
    assert res.ok, res.error
    assert any("eingefügt" in c.message for c in res.cells)
    _assert_checkboxes_follow_rows(TPL_DE, out)


def test_sync_uebernimmt_festtext_und_umbruch(tmp_path):
    import re
    import zipfile
    from spezismith.engine import apply_to_file, Rule
    tpl = "templates/vorlage_de.xlsx"
    spezi = str(tmp_path / "spezi.xlsx")
    apply_to_file(tpl, spezi, [Rule(coord="C78", new="nichts Vakuum", expected="Vakuum")])
    out = str(tmp_path / "out.xlsx")
    res = apply_to_file(spezi, out, [], template=tpl)
    assert res.ok
    import openpyxl
    assert openpyxl.load_workbook(out).worksheets[0]["C78"].value == "Vakuum"
    xml = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode()
    assert re.search(r'<brk id="43"', xml)
